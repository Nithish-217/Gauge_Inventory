import sys
from pathlib import Path
from typing import List, Any
from openpyxl import load_workbook


def first_non_empty_row(ws, max_scan: int = 50) -> int:
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1):
        if any(cell is not None and str(cell).strip() != "" for cell in row):
            return i
    return 1


def preview_sheet(ws, header_row: int, preview_rows: int = 5, max_cols: int = 20):
    # Gather headers
    headers: List[str] = []
    for cell in ws.iter_rows(min_row=header_row, max_row=header_row, max_col=max_cols, values_only=True):
        headers = [str(c).strip() if c is not None else "" for c in cell]
    print("Columns:", headers)

    # Data rows
    start = header_row + 1
    end = header_row + preview_rows
    for row in ws.iter_rows(min_row=start, max_row=end, max_col=max_cols, values_only=True):
        values: List[Any] = ["" if v is None else v for v in row]
        print(values)


def inspect_excel(path: Path, max_rows: int = 5, write_report: bool = True):
    wb = load_workbook(filename=path, data_only=True, read_only=True)
    lines: List[str] = []
    lines.append(f"Workbook: {path}")
    lines.append("Sheets: " + ", ".join(wb.sheetnames))
    for name in wb.sheetnames:
        ws = wb[name]
        lines.append(f"\n[Sheet: {name}] Preview:")
        header_row = first_non_empty_row(ws)
        # headers
        headers: List[str] = []
        for cell in ws.iter_rows(min_row=header_row, max_row=header_row, max_col=50, values_only=True):
            headers = ["" if c is None else str(c).strip() for c in cell]
        lines.append("Columns: " + ", ".join(headers))
        # rows
        start = header_row + 1
        end = header_row + max_rows
        for row in ws.iter_rows(min_row=start, max_row=end, max_col=50, values_only=True):
            values = ["" if v is None else str(v) for v in row]
            lines.append("[row] " + ", ".join(values))

    report_path = None
    if write_report:
        report_path = path.with_name(path.stem + "_inspect.txt")
        report_path.write_text("\n".join(lines), encoding="utf-8")

    # Also print concise
    print("\n".join(lines[:50]))
    if report_path:
        print(f"\n[Report written to] {report_path}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python inspect_excel.py <path-to-excel>")
        sys.exit(1)
    path = Path(sys.argv[1])
    if not path.exists():
        print(f"File not found: {path}")
        sys.exit(1)
    inspect_excel(path)
